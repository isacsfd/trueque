#!/usr/bin/env python3
"""Generate the physical reconciliation workbook from the local inventory DB."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
from urllib.parse import urlparse

import psycopg
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1", "postgres"}
BLUE = "174A91"
BLUE_DARK = "10386F"
BLUE_SOFT = "EDF3FB"
YELLOW = "FFD447"
YELLOW_SOFT = "FFF9E4"
WHITE = "FFFFFF"
INK = "172B46"
MUTED = "5D6D81"
BORDER = "DFE6EF"
GREEN = "1E7A45"
GREEN_SOFT = "E8F7EE"
AMBER = "B7791F"
AMBER_SOFT = "FFF4D6"
RED = "A93226"
RED_SOFT = "FDECEC"
GREY_SOFT = "F7F9FC"

THIN = Side(style="thin", color=BORDER)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def clean(value: object) -> str:
    return "" if value is None else " ".join(str(value).split())


def local_database_url() -> str:
    url = os.environ.get("DATABASE_MIGRATION_URL", "").strip()
    if not url:
        raise SystemExit("DATABASE_MIGRATION_URL is required.")
    host = (urlparse(url).hostname or "").lower()
    if host not in LOCAL_HOSTS:
        raise SystemExit(f"Refusing non-local database host: {host}")
    return url


def load_rows() -> tuple[list[dict[str, object]], str]:
    with psycopg.connect(local_database_url()) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT r.source_row, r.source_number, r.source_copy_number,
                       b.title, b.author, COALESCE(p.name, '') AS publisher,
                       c.inventory_code, c.location, c.bibliographic_record_id
                FROM public.inventory_import_rows r
                JOIN public.copies c ON c.id = r.copy_id
                JOIN public.bibliographic_records b ON b.id = c.bibliographic_record_id
                LEFT JOIN public.publishers p ON p.id = b.publisher_id
                WHERE r.status IN ('imported', 'reused')
                ORDER BY c.inventory_code
                """
            )
            columns = [description.name for description in cursor.description]
            rows = [dict(zip(columns, values)) for values in cursor.fetchall()]
    if not rows:
        raise SystemExit("No imported inventory rows found. Run the importer first.")
    return rows, hashlib.sha256(
        Path("INVENTARIO TRUEQUE.xlsx").read_bytes()
    ).hexdigest() if Path("INVENTARIO TRUEQUE.xlsx").exists() else ""


def header_row(ws, row: int, headers: list[str], height: int = 34) -> None:
    for index, value in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=index, value=value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = height


def title_block(ws, last_column: str, subtitle: str) -> None:
    ws.merge_cells(f"A1:{last_column}1")
    ws["A1"] = "CONCILIACIÓN FÍSICA DEL INVENTARIO"
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Aptos", size=18, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 38

    ws.merge_cells(f"A2:{last_column}2")
    ws["A2"] = subtitle
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE_SOFT)
    ws["A2"].font = Font(name="Aptos", size=10, color=BLUE_DARK)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 24

    ws.merge_cells(f"A3:{last_column}3")
    ws["A3"] = "Documento interno de verificación · No es una autorización de habilitación · No modifica el inventario maestro"
    ws["A3"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A3"].font = Font(name="Aptos", size=9, bold=True, color=BLUE_DARK)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 22


def add_validation(ws, formula: str, cell_range: str, title: str, prompt: str) -> None:
    validation = DataValidation(
        type="list",
        formula1=f'"{formula}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Valor no permitido",
        error="Selecciona uno de los valores de la lista.",
        promptTitle=title,
        prompt=prompt,
    )
    ws.add_data_validation(validation)
    validation.add(cell_range)


def add_date_validation(ws, cell_range: str) -> None:
    validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Fecha no válida",
        error="Introduce una fecha entre 2020 y 2100.",
    )
    ws.add_data_validation(validation)
    validation.add(cell_range)


def add_conditional(ws, cell_range: str, formulas: list[tuple[str, str, str]]) -> None:
    for formula, fill_color, font_color in formulas:
        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[formula],
                fill=PatternFill("solid", fgColor=fill_color),
                font=Font(color=font_color, bold=True),
            ),
        )


def build_conciliacion(wb: Workbook, rows: list[dict[str, object]]) -> int:
    ws = wb.active
    ws.title = "Conciliación física"
    ws.sheet_properties.tabColor = BLUE
    headers = [
        "Código institucional",
        "Número fuente",
        "Título",
        "Autor",
        "Editorial",
        "Número / identificación del ejemplar",
        "Ubicación registrada",
        "Presencia física",
        "Estado físico",
        "Identificación bibliográfica",
        "Disponibilidad propuesta",
        "Discrepancia encontrada",
        "Acción recomendada",
        "Responsable de verificación",
        "Fecha de verificación",
        "Observaciones",
    ]
    title_block(ws, "P", "Sede Plaza de la Paz · Inventario institucional · Trueque IUB")
    header_row(ws, 4, headers)
    first_data_row = 5
    last_data_row = first_data_row + len(rows) - 1

    for offset, row in enumerate(rows):
        excel_row = first_data_row + offset
        values = [
            clean(row["inventory_code"]),
            clean(row["source_number"]),
            clean(row["title"]),
            clean(row["author"]),
            clean(row["publisher"]),
            clean(row["source_copy_number"]),
            clean(row["location"]),
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=column, value=value)
            cell.border = BORDER_ALL
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in (3, 4, 5, 12, 13, 16),
            )
        ws.row_dimensions[excel_row].height = 32

    table = Table(
        displayName="TablaConciliacion",
        ref=f"A4:P{last_data_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    widths = {
        "A": 18, "B": 10, "C": 36, "D": 27, "E": 23, "F": 15,
        "G": 21, "H": 15, "I": 14, "J": 20, "K": 20, "L": 38,
        "M": 38, "N": 22, "O": 17, "P": 30,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:P{last_data_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    add_validation(ws, "Pendiente,Sí,No", f"H5:H{last_data_row}", "Presencia física", "Pendiente, Sí o No.")
    add_validation(ws, "Pendiente,Bueno,Regular,Malo", f"I5:I{last_data_row}", "Estado físico", "Pendiente, Bueno, Regular o Malo.")
    add_validation(ws, "Pendiente,Coincide,No coincide", f"J5:J{last_data_row}", "Identificación bibliográfica", "Pendiente, Coincide o No coincide.")
    add_validation(ws, "Pendiente,Habilitar,Bloquear,Revisar", f"K5:K{last_data_row}", "Disponibilidad propuesta", "Propuesta sujeta a aprobación institucional.")
    add_date_validation(ws, f"O5:O{last_data_row}")
    for excel_row in range(first_data_row, last_data_row + 1):
        ws.cell(row=excel_row, column=15).number_format = "DD/MM/YYYY"
        ws.cell(row=excel_row, column=15).alignment = Alignment(horizontal="center", vertical="top")

    add_conditional(ws, f"H5:H{last_data_row}", [
        ("$H5=\"No\"", RED_SOFT, RED),
        ("$H5=\"Sí\"", GREEN_SOFT, GREEN),
        ("$H5=\"Pendiente\"", AMBER_SOFT, AMBER),
        ("LEN(TRIM($H5))=0", AMBER_SOFT, AMBER),
    ])
    add_conditional(ws, f"I5:I{last_data_row}", [
        ("$I5=\"Malo\"", RED_SOFT, RED),
        ("$I5=\"Regular\"", AMBER_SOFT, AMBER),
        ("$I5=\"Bueno\"", GREEN_SOFT, GREEN),
        ("$I5=\"Pendiente\"", AMBER_SOFT, AMBER),
        ("LEN(TRIM($I5))=0", AMBER_SOFT, AMBER),
    ])
    add_conditional(ws, f"J5:J{last_data_row}", [
        ("$J5=\"No coincide\"", RED_SOFT, RED),
        ("$J5=\"Coincide\"", GREEN_SOFT, GREEN),
        ("$J5=\"Pendiente\"", AMBER_SOFT, AMBER),
        ("LEN(TRIM($J5))=0", AMBER_SOFT, AMBER),
    ])
    add_conditional(ws, f"K5:K{last_data_row}", [
        ("$K5=\"Bloquear\"", RED_SOFT, RED),
        ("$K5=\"Habilitar\"", GREEN_SOFT, GREEN),
        ("$K5=\"Revisar\"", AMBER_SOFT, AMBER),
        ("$K5=\"Pendiente\"", AMBER_SOFT, AMBER),
        ("LEN(TRIM($K5))=0", AMBER_SOFT, AMBER),
    ])
    add_conditional(ws, f"L5:L{last_data_row}", [
        ("LEN(TRIM($L5))>0", RED_SOFT, RED),
    ])
    return last_data_row


def build_summary(wb: Workbook, last_data_row: int, source_hash: str) -> None:
    ws = wb.create_sheet("Resumen")
    ws.sheet_properties.tabColor = YELLOW
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:C1")
    ws["A1"] = "RESUMEN DE CONCILIACIÓN"
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Aptos", size=18, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:C2")
    ws["A2"] = "Indicadores calculados sobre las 67 filas. Los campos vacíos o “Pendiente” no cuentan como inspección."
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE_SOFT)
    ws["A2"].font = Font(name="Aptos", size=10, color=BLUE_DARK)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 34

    metrics = [
        ("Total de ejemplares registrados", f"=COUNTA('Conciliación física'!A5:A{last_data_row})"),
        ("Ejemplares inspeccionados", f"=COUNTIF('Conciliación física'!H5:H{last_data_row},\"Sí\")+COUNTIF('Conciliación física'!H5:H{last_data_row},\"No\")"),
        ("Ejemplares pendientes", "=B4-B5"),
        ("Ejemplares encontrados", f"=COUNTIF('Conciliación física'!H5:H{last_data_row},\"Sí\")"),
        ("Ejemplares no encontrados", f"=COUNTIF('Conciliación física'!H5:H{last_data_row},\"No\")"),
        ("Ejemplares con identificación coincidente", f"=COUNTIF('Conciliación física'!J5:J{last_data_row},\"Coincide\")"),
        ("Ejemplares con discrepancias", f"=SUMPRODUCT(--(('Conciliación física'!H5:H{last_data_row}=\"No\")+('Conciliación física'!J5:J{last_data_row}=\"No coincide\")+('Conciliación física'!L5:L{last_data_row}<>\"\")>0))"),
        ("Ejemplares en buen estado", f"=COUNTIF('Conciliación física'!I5:I{last_data_row},\"Bueno\")"),
        ("Ejemplares en estado regular", f"=COUNTIF('Conciliación física'!I5:I{last_data_row},\"Regular\")"),
        ("Ejemplares en mal estado", f"=COUNTIF('Conciliación física'!I5:I{last_data_row},\"Malo\")"),
        ("Propuestas de habilitación", f"=COUNTIF('Conciliación física'!K5:K{last_data_row},\"Habilitar\")"),
        ("Propuestas de bloqueo", f"=COUNTIF('Conciliación física'!K5:K{last_data_row},\"Bloquear\")"),
        ("Casos pendientes de decisión", f"=COUNTBLANK('Conciliación física'!K5:K{last_data_row})+COUNTIF('Conciliación física'!K5:K{last_data_row},\"Pendiente\")+COUNTIF('Conciliación física'!K5:K{last_data_row},\"Revisar\")"),
        ("Incidencias abiertas o en revisión", "=COUNTIF(Incidencias!G5:G24,\"Abierta\")+COUNTIF(Incidencias!G5:G24,\"En revisión\")"),
    ]
    ws["A3"] = "Indicador"
    ws["B3"] = "Valor"
    for column in ("A", "B"):
        ws[f"{column}3"].fill = PatternFill("solid", fgColor=BLUE)
        ws[f"{column}3"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        ws[f"{column}3"].border = BORDER_ALL
    for index, (label, formula) in enumerate(metrics, start=4):
        ws[f"A{index}"] = label
        ws[f"B{index}"] = formula
        ws[f"A{index}"].font = Font(name="Aptos", size=10, color=INK)
        ws[f"B{index}"].font = Font(name="Aptos", size=11, bold=True, color=BLUE_DARK)
        ws[f"A{index}"].border = BORDER_ALL
        ws[f"B{index}"].border = BORDER_ALL
        ws[f"B{index}"].alignment = Alignment(horizontal="center")
        ws[f"B{index}"].number_format = "0"
        ws.row_dimensions[index].height = 24

    ws.merge_cells("A20:C20")
    ws["A20"] = "Regla: “Pendiente” y las celdas vacías no son resultados confirmados. “Habilitar” es una propuesta, no una autorización."
    ws["A20"].fill = PatternFill("solid", fgColor=YELLOW_SOFT)
    ws["A20"].font = Font(name="Aptos", size=10, bold=True, color=BLUE_DARK)
    ws["A20"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[20].height = 42

    ws["A22"] = "Fuente"
    ws["B22"] = "INVENTARIO TRUEQUE.xlsx · Hoja1"
    ws["A23"] = "SHA-256"
    ws["B23"] = source_hash or "No disponible"
    ws["A24"] = "Base local"
    ws["B24"] = "iub_local · importación 0006"
    ws["A25"] = "Nota"
    ws["B25"] = "Una habilitación definitiva requiere aprobación institucional y no se aplica desde este archivo."
    for row in range(22, 26):
        ws[f"A{row}"].font = Font(name="Aptos", size=9, bold=True, color=MUTED)
        ws[f"B{row}"].font = Font(name="Aptos", size=9, color=MUTED)
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 24
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 45
    ws.freeze_panes = "A4"


def build_incidencias(wb: Workbook) -> None:
    ws = wb.create_sheet("Incidencias")
    ws.sheet_properties.tabColor = "A93226"
    ws.sheet_view.showGridLines = False
    title_block(ws, "I", "Registro de hallazgos, evidencias, responsables y resolución aprobada")
    headers = [
        "Código institucional",
        "Tipo de incidencia",
        "Descripción",
        "Evidencia / referencia",
        "Responsable",
        "Fecha",
        "Estado de resolución",
        "Resolución aprobada",
        "Observaciones",
    ]
    header_row(ws, 4, headers)
    for excel_row in range(5, 25):
        for column in range(1, 10):
            cell = ws.cell(row=excel_row, column=column)
            cell.border = BORDER_ALL
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 8, 9))
        ws.row_dimensions[excel_row].height = 28
    table = Table(displayName="TablaIncidencias", ref="A4:I24")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.freeze_panes = "A5"
    add_validation(ws, "Presencia física,Estado físico,Identificación bibliográfica,Disponibilidad,Datos maestros,Otro", "B5:B24", "Tipo de incidencia", "Clasifica el hallazgo.")
    add_validation(ws, "Abierta,En revisión,Resuelta,Rechazada", "G5:G24", "Estado de resolución", "Abierta, En revisión, Resuelta o Rechazada.")
    add_date_validation(ws, "F5:F24")
    for excel_row in range(5, 25):
        ws.cell(row=excel_row, column=6).number_format = "DD/MM/YYYY"
    widths = {"A": 18, "B": 28, "C": 42, "D": 30, "E": 22, "F": 16, "G": 22, "H": 35, "I": 28}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    add_conditional(ws, "G5:G24", [
        ("$G5=\"Resuelta\"", GREEN_SOFT, GREEN),
        ("$G5=\"Abierta\"", RED_SOFT, RED),
        ("$G5=\"En revisión\"", AMBER_SOFT, AMBER),
        ("LEN(TRIM($G5))=0", GREY_SOFT, MUTED),
    ])
    ws.print_title_rows = "1:4"
    ws.print_area = "A1:I24"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_instructions(wb: Workbook, total: int, source_hash: str) -> None:
    ws = wb.create_sheet("Instrucciones")
    ws.sheet_properties.tabColor = BLUE_DARK
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "INSTRUCCIONES DE CONCILIACIÓN FÍSICA"
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Aptos", size=18, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:F2")
    ws["A2"] = f"Trueque IUB · Sede Plaza de la Paz · {total} ejemplares precargados · documento interno"
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE_SOFT)
    ws["A2"].font = Font(name="Aptos", size=10, color=BLUE_DARK)
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 26

    sections = [
        ("1. Objetivo", [
            "Verificar físicamente cada ejemplar institucional y documentar si existe, su estado, la coincidencia bibliográfica y una propuesta de disponibilidad.",
            "La plantilla no modifica el inventario maestro, no cambia la base de datos y no autoriza por sí sola la habilitación en el catálogo.",
        ]),
        ("2. Cómo usarla", [
            "Trabaja en la hoja “Conciliación física”, una fila por ejemplar. No elimines ni reordenes ejemplares; los títulos repetidos representan ejemplares distintos.",
            "Completa Presencia física, Estado físico, Identificación bibliográfica y Disponibilidad propuesta usando las listas desplegables.",
            "Registra Responsable y Fecha cuando marques un resultado. Si hay una discrepancia, usa la hoja “Incidencias” y documenta evidencia.",
            "No marques “Sí”, “Coincide” o “Habilitar” sin verificación humana. “Pendiente” y las celdas vacías no cuentan como inspección.",
        ]),
        ("3. Estados", [
            "Presencia: Pendiente / Sí / No.",
            "Estado: Pendiente / Bueno / Regular / Malo.",
            "Identificación: Pendiente / Coincide / No coincide.",
            "Disponibilidad propuesta: Pendiente / Habilitar / Bloquear / Revisar.",
            "Habilitar es una recomendación; Bloquear y Revisar requieren acción institucional.",
        ]),
        ("4. Datos de origen", [
            "Código institucional: IUB-PAZ-001 a IUB-PAZ-067, derivado del número fuente del Excel.",
            "Se conservan títulos, autores, editoriales, números de ejemplar y la ubicación PLAZA DE LA PAZ.",
            "Los valores N/A y N/E se conservan como aparecen en la fuente; una corrección solo puede registrarse como incidencia aprobada.",
            "No se inventaron ISBN. Los registros bibliográficos no son ejemplares físicos.",
        ]),
        ("5. Cierre", [
            "Una propuesta de habilitación no es una autorización. La habilitación definitiva requiere aprobación institucional y una actualización controlada de la base.",
            "Conserva este archivo como evidencia de la conciliación. No sobrescribas el inventario original.",
        ]),
    ]
    row = 4
    for heading, bullets in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=heading)
        cell.font = Font(name="Aptos", size=11, bold=True, color=BLUE_DARK)
        cell.fill = PatternFill("solid", fgColor=YELLOW_SOFT)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 26
        row += 1
        for bullet in bullets:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell = ws.cell(row=row, column=1, value=f"• {bullet}")
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = 34
            row += 1
        row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f"Fuente: INVENTARIO TRUEQUE.xlsx · Hoja1 · SHA-256: {source_hash or 'no disponible'}")
    ws.cell(row=row, column=1).font = Font(name="Aptos", size=8, color=MUTED)
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    for column in ("B", "C", "D", "E", "F"):
        ws.column_dimensions[column].width = 12


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        default="docs/inventario/Conciliacion_Fisica_Trueque_IUB.xlsx",
    )
    args = parser.parse_args()

    rows, source_hash = load_rows()
    codes = [clean(row["inventory_code"]) for row in rows]
    if len(rows) != 67:
        raise SystemExit(f"Expected 67 imported copies, found {len(rows)}.")
    if len(set(codes)) != len(codes):
        raise SystemExit("Duplicate institutional codes found in the source data.")
    if len({row["bibliographic_record_id"] for row in rows}) != 65:
        raise SystemExit("Expected 65 distinct bibliographic records.")

    wb = Workbook()
    last_data_row = build_conciliacion(wb, rows)
    build_summary(wb, last_data_row, source_hash)
    build_incidencias(wb)
    build_instructions(wb, len(rows), source_hash)
    wb.properties.title = "Conciliación física del inventario institucional · Trueque IUB"
    wb.properties.creator = "Trueque IUB"
    wb.properties.subject = "Verificación física de ejemplares · Plaza de la Paz"
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(json.dumps({
        "output": str(output.resolve()),
        "rows": len(rows),
        "bibliographic_records": len({row["bibliographic_record_id"] for row in rows}),
        "codes": len(set(codes)),
        "source_sha256": source_hash,
        "sheets": wb.sheetnames,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
