#!/usr/bin/env python3
"""Process a completed physical inventory reconciliation workbook.

The processor is read-only for both the source workbook and PostgreSQL. It
produces a separate report and never changes availability, moves inventory,
resolves incidents, or enables copies.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

PROCESSOR_VERSION = "1.0.1"
LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1", "postgres"}
INPUT_HEADERS = [
    "Código institucional",
    "Número fuente",
    "Título",
    "Autor",
    "Editorial",
    "Número / identificación del ejemplar",
    "Ubicación registrada",
    "Presencia física",
    "Estado físico",
    "Identificación bibliográfica",
    "Disponibilidad propuesta",
    "Discrepancia encontrada",
    "Acción recomendada",
    "Responsable de verificación",
    "Fecha de verificación",
    "Observaciones",
]
DETAIL_HEADERS = [
    "Código institucional",
    "Número fuente",
    "Título registrado",
    "Autor registrado",
    "Editorial registrada",
    "Ejemplar",
    "Ubicación registrada",
    "Presencia física",
    "Estado físico",
    "Identificación bibliográfica",
    "Disponibilidad propuesta",
    "Discrepancia declarada en el archivo",
    "Acción recomendada en el archivo",
    "Resultado físico",
    "Datos registrados coinciden",
    "Clasificación",
    "Motivo de la decisión",
    "Incidencias abiertas",
    "Discrepancias detectadas",
    "Responsable de verificación",
    "Fecha de verificación",
    "Observaciones",
]
INCIDENT_HEADERS = [
    "Código institucional",
    "Tipo de incidencia",
    "Descripción",
    "Evidencia / referencia",
    "Responsable",
    "Fecha",
    "Estado de resolución",
    "Resolución aprobada",
    "Observaciones",
]
INCIDENT_TYPES = {
    "Presencia física",
    "Estado físico",
    "Identificación bibliográfica",
    "Disponibilidad",
    "Datos maestros",
    "Otro",
}
INCIDENT_STATUSES = {"Abierta", "En revisión", "Resuelta", "Rechazada"}
PRESENCE_VALUES = {"Pendiente", "Sí", "No"}
CONDITION_VALUES = {"Pendiente", "Bueno", "Regular", "Malo"}
IDENTIFICATION_VALUES = {"Pendiente", "Coincide", "No coincide"}
AVAILABILITY_VALUES = {"Pendiente", "Habilitar", "Bloquear", "Revisar"}

BLUE = "174A91"
BLUE_DARK = "10386F"
BLUE_SOFT = "EDF3FB"
YELLOW = "FFD447"
YELLOW_SOFT = "FFF9E4"
WHITE = "FFFFFF"
INK = "172B46"
MUTED = "5D6D81"
BORDER = "DFE6EF"
GREEN = "1E7A45"
GREEN_SOFT = "E8F7EE"
AMBER = "B7791F"
AMBER_SOFT = "FFF4D6"
RED = "A93226"
RED_SOFT = "FDECEC"
GREY_SOFT = "F7F9FC"

THIN = Side(style="thin", color=BORDER)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def clean(value: object) -> str:
    return "" if value is None else " ".join(str(value).split())


def normalize(value: object) -> str:
    return clean(value).casefold()


def parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"invalid date: {text}")


def local_database_url() -> str:
    url = os.environ.get("DATABASE_MIGRATION_URL", "").strip()
    if not url:
        raise SystemExit("DATABASE_MIGRATION_URL is required.")
    host = (urlparse(url).hostname or "").lower()
    if host not in LOCAL_HOSTS:
        raise SystemExit(f"Refusing non-local database host: {host}")
    return url


@dataclass
class ReferenceCopy:
    code: str
    source_number: str
    copy_number: str
    title: str
    author: str
    publisher: str
    location: str
    availability: str
    provenance_type: str
    verified: bool
    bibliographic_id: int
    source_row: int


@dataclass
class Issue:
    code: str
    incident_type: str
    description: str
    evidence: str
    responsible: str
    date_value: date | None
    status: str
    resolution: str
    observations: str
    blocking: bool = False


@dataclass
class RowResult:
    code: str
    source_row: int
    reference: ReferenceCopy | None
    presence: str
    condition: str
    identification: str
    availability: str
    declared_discrepancy: str
    declared_action: str
    result_physical: str
    data_matches: bool
    classification: str
    reason: str
    open_incidents: int
    discrepancies: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    responsible: str = ""
    date_value: date | None = None
    observations: str = ""
    candidate: bool = False
    no_habilitable: bool = False
    confirmed: bool = False


def load_reference(url: str) -> dict[str, ReferenceCopy]:
    with psycopg.connect(url) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT c.inventory_code, r.source_number, r.source_copy_number,
                       b.title, b.author, COALESCE(p.name, ''),
                       c.location, c.availability, c.provenance_type,
                       (c.verified_at IS NOT NULL), b.id, r.source_row
                FROM public.inventory_import_rows r
                JOIN public.copies c ON c.id = r.copy_id
                JOIN public.bibliographic_records b ON b.id = c.bibliographic_record_id
                LEFT JOIN public.publishers p ON p.id = b.publisher_id
                WHERE r.status IN ('imported', 'reused')
                ORDER BY c.inventory_code
                """
            )
            rows = cursor.fetchall()
    return {
        row[0]: ReferenceCopy(
            code=row[0],
            source_number=clean(row[1]),
            copy_number=clean(row[2]),
            title=clean(row[3]),
            author=clean(row[4]),
            publisher=clean(row[5]),
            location=clean(row[6]),
            availability=clean(row[7]),
            provenance_type=clean(row[8]),
            verified=bool(row[9]),
            bibliographic_id=int(row[10]),
            source_row=int(row[11]),
        )
        for row in rows
    }


def read_incidents(workbook: Workbook, expected_codes: set[str]) -> tuple[list[Issue], list[str]]:
    if "Incidencias" not in workbook.sheetnames:
        return [], ["Falta la hoja Incidencias."]
    ws = workbook["Incidencias"]
    headers = [clean(ws.cell(row=4, column=column).value) for column in range(1, 10)]
    if headers != INCIDENT_HEADERS:
        return [], [f"Encabezados de Incidencias inesperados: {headers!r}"]
    issues: list[Issue] = []
    global_errors: list[str] = []
    for source_row in range(5, ws.max_row + 1):
        values = [ws.cell(row=source_row, column=column).value for column in range(1, 10)]
        if all(clean(value) == "" for value in values):
            continue
        code = clean(values[0])
        incident_type = clean(values[1])
        status = clean(values[6])
        resolution = clean(values[7])
        try:
            issue_date = parse_date(values[5])
        except ValueError as error:
            global_errors.append(f"Incidencia fila {source_row}: {error}")
            issue_date = None
        if code not in expected_codes:
            global_errors.append(f"Incidencia fila {source_row}: código desconocido {code!r}.")
        if incident_type not in INCIDENT_TYPES:
            global_errors.append(f"Incidencia fila {source_row}: tipo inválido {incident_type!r}.")
        if status not in INCIDENT_STATUSES:
            global_errors.append(f"Incidencia fila {source_row}: estado inválido {status!r}.")
        blocking = status in {"Abierta", "En revisión"} or (status == "Resuelta" and not resolution)
        if status == "Resuelta" and not resolution:
            global_errors.append(f"Incidencia fila {source_row}: resuelta sin resolución aprobada.")
        issues.append(Issue(
            code=code,
            incident_type=incident_type,
            description=clean(values[2]),
            evidence=clean(values[3]),
            responsible=clean(values[4]),
            date_value=issue_date,
            status=status,
            resolution=resolution,
            observations=clean(values[8]),
            blocking=blocking,
        ))
    return issues, global_errors


def read_detail(workbook: Workbook, reference: dict[str, ReferenceCopy]) -> tuple[list[RowResult], list[str], list[Issue]]:
    if "Conciliación física" not in workbook.sheetnames:
        raise SystemExit("Falta la hoja Conciliación física.")
    ws = workbook["Conciliación física"]
    headers = [clean(ws.cell(row=4, column=column).value) for column in range(1, 17)]
    if headers != INPUT_HEADERS:
        raise SystemExit(f"Encabezados de Conciliación física inesperados: {headers!r}")

    global_errors: list[str] = []
    issues, incident_errors = read_incidents(workbook, set(reference))
    global_errors.extend(incident_errors)
    open_by_code: dict[str, int] = {}
    blocking_by_code: dict[str, int] = {}
    for issue in issues:
        if issue.code:
            if issue.blocking:
                blocking_by_code[issue.code] = blocking_by_code.get(issue.code, 0) + 1
            if issue.status in {"Abierta", "En revisión"}:
                open_by_code[issue.code] = open_by_code.get(issue.code, 0) + 1

    raw_rows: list[tuple[int, list[object]]] = []
    for source_row in range(5, ws.max_row + 1):
        values = [ws.cell(row=source_row, column=column).value for column in range(1, 17)]
        if all(clean(value) == "" for value in values):
            continue
        raw_rows.append((source_row, values))

    code_rows: dict[str, list[tuple[int, list[object]]]] = {}
    for source_row, values in raw_rows:
        code = clean(values[0])
        if code:
            code_rows.setdefault(code, []).append((source_row, values))

    results: list[RowResult] = []
    for source_row, values in raw_rows:
        code = clean(values[0])
        file_title = clean(values[2])
        file_author = clean(values[3])
        file_publisher = clean(values[4])
        file_copy = clean(values[5])
        file_location = clean(values[6])
        presence = clean(values[7]) or "Pendiente"
        condition = clean(values[8]) or "Pendiente"
        identification = clean(values[9]) or "Pendiente"
        availability = clean(values[10]) or "Pendiente"
        declared_discrepancy = clean(values[11])
        declared_action = clean(values[12])
        responsible = clean(values[13])
        observations = clean(values[15])
        try:
            verification_date = parse_date(values[14])
        except ValueError as error:
            verification_date = None
            global_errors.append(f"Fila {source_row}: {error}")

        reference_row = reference.get(code)
        errors: list[str] = []
        discrepancies: list[str] = []
        if not code:
            errors.append("Código institucional ausente.")
        elif reference_row is None:
            errors.append("Código institucional desconocido.")
        else:
            if len(code_rows.get(code, [])) > 1:
                errors.append("Código institucional duplicado en el Excel.")
            if normalize(file_title) != normalize(reference_row.title):
                discrepancies.append(f"Título del archivo no coincide: {file_title!r}.")
            if normalize(file_author) != normalize(reference_row.author):
                discrepancies.append(f"Autor del archivo no coincide: {file_author!r}.")
            if normalize(file_publisher) != normalize(reference_row.publisher):
                discrepancies.append(f"Editorial del archivo no coincide: {file_publisher!r}.")
            if normalize(file_copy) != normalize(reference_row.copy_number):
                discrepancies.append(f"Ejemplar del archivo no coincide: {file_copy!r}.")
            if normalize(file_location) != normalize(reference_row.location):
                discrepancies.append(f"Ubicación del archivo no coincide: {file_location!r}.")
            if reference_row.provenance_type == "institutional" and not reference_row.verified:
                if reference_row.availability not in ("unavailable", "maintenance", "retired"):
                    discrepancies.append(
                        f"Ejemplar institucional no verificado con disponibilidad {reference_row.availability}."
                    )
            elif reference_row.availability != "available":
                discrepancies.append(f"Ejemplar no disponible en la base: {reference_row.availability}.")
            if reference_row.provenance_type != "institutional":
                discrepancies.append(f"Procedencia no institucional: {reference_row.provenance_type}.")

        if presence not in PRESENCE_VALUES:
            errors.append(f"Presencia inválida: {presence!r}.")
        if condition not in CONDITION_VALUES:
            errors.append(f"Estado físico inválido: {condition!r}.")
        if identification not in IDENTIFICATION_VALUES:
            errors.append(f"Identificación inválida: {identification!r}.")
        if availability not in AVAILABILITY_VALUES:
            errors.append(f"Disponibilidad inválida: {availability!r}.")
        if verification_date and not (date(2020, 1, 1) <= verification_date <= date(2100, 12, 31)):
            errors.append("Fecha de verificación fuera de rango.")
        if presence in {"Sí", "No"} and not responsible:
            errors.append("Resultado físico sin responsable.")
        if presence in {"Sí", "No"} and verification_date is None:
            errors.append("Resultado físico sin fecha de verificación.")

        if errors or reference_row is None:
            result_physical = "No válido"
        elif presence == "Sí":
            result_physical = "Confirmado"
        elif presence == "No":
            result_physical = "No encontrado"
        else:
            result_physical = "Pendiente"
        open_count = open_by_code.get(code, 0)
        blocking_count = blocking_by_code.get(code, 0)
        has_discrepancy = bool(discrepancies) or bool(declared_discrepancy) or identification == "No coincide" or blocking_count > 0
        data_matches = not discrepancies
        candidate = False
        no_habilitable = False
        reason_parts: list[str] = []

        if errors:
            classification = "Error de validación"
            reason_parts.append("; ".join(errors))
        elif reference_row is None:
            classification = "Error de validación"
            reason_parts.append("El código no pertenece al inventario institucional local.")
        elif presence == "No":
            classification = "No encontrado"
            reason_parts.append("El conteo físico indica que el ejemplar no fue encontrado.")
            no_habilitable = True
        elif identification == "No coincide":
            classification = "Identificación no coincidente"
            reason_parts.append("La identificación bibliográfica no coincide con el registro.")
            no_habilitable = True
        elif has_discrepancy:
            classification = "Con discrepancia"
            reason_parts.append("; ".join(discrepancies) if discrepancies else "Hay incidencias o discrepancias documentadas.")
            if declared_discrepancy:
                reason_parts.append(f"Discrepancia declarada: {declared_discrepancy}")
            if blocking_count:
                reason_parts.append(f"{blocking_count} incidencia(s) bloqueante(s) sin resolver.")
        elif presence in {"", "Pendiente"}:
            classification = "Pendiente de conteo"
            reason_parts.append("No hay resultado físico confirmado.")
        elif condition in {"", "Pendiente"}:
            classification = "Estado físico pendiente"
            reason_parts.append("La presencia está confirmada, pero falta evaluar el estado.")
        elif condition in {"Regular", "Malo"}:
            classification = "Pendiente de decisión administrativa"
            reason_parts.append(
                f"Estado físico {condition}: el umbral de aceptación D-02 está pendiente; "
                "no se propone automáticamente como candidato."
            )
        elif availability == "Bloquear":
            classification = "No habilitable con la información disponible"
            reason_parts.append("La disponibilidad propuesta es Bloquear.")
            no_habilitable = True
        elif availability == "Habilitar":
            candidate = True
            classification = "Candidato a habilitación"
            reason_parts.append("Presencia, identificación, estado y datos institucionales conciliados; propuesta explícita de habilitar.")
            reason_parts.append("Requiere aprobación institucional; no existe una política de aceptación de estado en la base.")
        else:
            classification = "Pendiente de decisión administrativa"
            reason_parts.append("Falta una propuesta de disponibilidad o está marcada para revisión.")

        if reference_row is not None:
            if reference_row.provenance_type == "institutional" and not reference_row.verified:
                pass
            elif reference_row.availability != "available":
                no_habilitable = True
        if reference_row is not None and reference_row.provenance_type != "institutional":
            no_habilitable = True
        if errors:
            no_habilitable = True

        results.append(RowResult(
            code=code or f"(fila {source_row})",
            source_row=source_row,
            reference=reference_row,
            presence=presence,
            condition=condition,
            identification=identification,
            availability=availability,
            declared_discrepancy=declared_discrepancy,
            declared_action=declared_action,
            result_physical=result_physical,
            data_matches=data_matches,
            classification=classification,
            reason=" ".join(reason_parts),
            open_incidents=open_count,
            discrepancies=discrepancies,
            errors=errors,
            responsible=responsible,
            date_value=verification_date,
            observations=observations,
            candidate=candidate,
            no_habilitable=no_habilitable,
            confirmed=presence == "Sí" and not errors and reference_row is not None,
        ))

    # Missing expected codes are reported without inventing file results.
    for code, ref in reference.items():
        if code in code_rows:
            continue
        global_errors.append(f"Código esperado ausente del Excel: {code}.")
        results.append(RowResult(
            code=code,
            source_row=0,
            reference=ref,
            presence="Pendiente",
            condition="Pendiente",
            identification="Pendiente",
            availability="Pendiente",
            declared_discrepancy="",
            declared_action="",
            result_physical="Pendiente",
            data_matches=True,
            classification="Error de validación",
            reason="Código esperado no aparece en el Excel; no se procesó ninguna fila física.",
            open_incidents=open_by_code.get(code, 0),
            errors=["Código esperado ausente del Excel."],
            no_habilitable=True,
        ))

    return results, global_errors, issues


def classify_summary(results: list[RowResult], global_errors: list[str], expected: int) -> dict[str, int]:
    row_errors = sum(1 for row in results if row.classification == "Error de validación")
    incident_errors = sum(
        1 for error in global_errors
        if error.startswith(("Incidencia", "Falta", "Encabezados"))
    )
    return {
        "total_expected": expected,
        "processed": sum(1 for row in results if row.source_row > 0),
        "confirmed": sum(1 for row in results if row.confirmed),
        "not_found": sum(1 for row in results if row.classification == "No encontrado"),
        "pending": sum(1 for row in results if row.classification in {"Pendiente de conteo", "Pendiente de decisión administrativa", "Estado físico pendiente"}),
        "discrepancies": sum(1 for row in results if row.classification == "Con discrepancia"),
        "candidates": sum(1 for row in results if row.candidate),
        "not_habilitable": sum(1 for row in results if row.no_habilitable),
        "validation_errors": row_errors + incident_errors,
        "incidents": 0,
    }


def header_row(ws, row: int, headers: list[str]) -> None:
    for column, value in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 36


def title_block(ws, last_column: str, subtitle: str) -> None:
    ws.merge_cells(f"A1:{last_column}1")
    ws["A1"] = "RESULTADO DE CONCILIACIÓN FÍSICA"
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].font = Font(name="Aptos", size=18, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells(f"A2:{last_column}2")
    ws["A2"] = subtitle
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE_SOFT)
    ws["A2"].font = Font(name="Aptos", size=10, color=BLUE_DARK)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 24
    ws.merge_cells(f"A3:{last_column}3")
    ws["A3"] = "Informe de lectura · No modifica inventario, disponibilidad, incidencias ni autorizaciones"
    ws["A3"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A3"].font = Font(name="Aptos", size=9, bold=True, color=BLUE_DARK)
    ws["A3"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[3].height = 22


def build_report(
    results: list[RowResult],
    issues: list[Issue],
    global_errors: list[str],
    reference_count: int,
    source_hash: str,
    processed_at: str,
    output: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_view.showGridLines = False
    title_block(ws, "B", f"Versión {PROCESSOR_VERSION} · {processed_at} · Fuente: {source_hash[:16]}…")

    summary = classify_summary(results, global_errors, reference_count)
    summary["incidents"] = len(issues)
    metrics = [
        ("Total de ejemplares esperados", summary["total_expected"]),
        ("Ejemplares procesados", summary["processed"]),
        ("Confirmados físicamente", summary["confirmed"]),
        ("No encontrados", summary["not_found"]),
        ("Pendientes", summary["pending"]),
        ("Con discrepancias", summary["discrepancies"]),
        ("Candidatos a habilitación", summary["candidates"]),
        ("No habilitables", summary["not_habilitable"]),
        ("Errores de validación", summary["validation_errors"]),
        ("Incidencias registradas", summary["incidents"]),
    ]
    ws["A5"] = "Indicador"
    ws["B5"] = "Valor"
    for column in ("A", "B"):
        ws[f"{column}5"].fill = PatternFill("solid", fgColor=BLUE)
        ws[f"{column}5"].font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        ws[f"{column}5"].border = BORDER_ALL
    for index, (label, value) in enumerate(metrics, start=6):
        ws[f"A{index}"] = label
        ws[f"B{index}"] = value
        ws[f"A{index}"].border = BORDER_ALL
        ws[f"B{index}"].border = BORDER_ALL
        ws[f"A{index}"].font = Font(name="Aptos", size=10, color=INK)
        ws[f"B{index}"].font = Font(name="Aptos", size=11, bold=True, color=BLUE_DARK)
        ws[f"B{index}"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[index].height = 24

    ws.merge_cells("A18:B18")
    ws["A18"] = "Un candidato es una propuesta sujeta a aprobación. Este informe no habilita ejemplares ni cambia la base."
    ws["A18"].fill = PatternFill("solid", fgColor=YELLOW_SOFT)
    ws["A18"].font = Font(name="Aptos", size=10, bold=True, color=BLUE_DARK)
    ws["A18"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[18].height = 42
    ws["A20"] = "Hash del Excel"
    ws["B20"] = source_hash
    ws["A21"] = "Procesador"
    ws["B21"] = f"Inventory reconciliation v{PROCESSOR_VERSION}"
    ws["A22"] = "Base de referencia"
    ws["B22"] = f"{reference_count} ejemplares institucionales locales"
    for row in range(20, 23):
        ws[f"A{row}"].font = Font(name="Aptos", size=9, bold=True, color=MUTED)
        ws[f"B{row}"].font = Font(name="Aptos", size=9, color=MUTED)
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 62

    detail = wb.create_sheet("Detalle")
    detail.sheet_properties.tabColor = BLUE_DARK
    detail.sheet_view.showGridLines = False
    title_block(detail, "V", "Detalle por ejemplar · La clasificación no autoriza cambios")
    header_row(detail, 4, DETAIL_HEADERS)
    for index, row in enumerate(results, start=5):
        ref = row.reference
        values = [
            row.code,
            ref.source_number if ref else "",
            ref.title if ref else "",
            ref.author if ref else "",
            ref.publisher if ref else "",
            ref.copy_number if ref else "",
            ref.location if ref else "",
            row.presence,
            row.condition,
            row.identification,
            row.availability,
            row.declared_discrepancy,
            row.declared_action,
            row.result_physical,
            "Sí" if row.data_matches else "No",
            row.classification,
            row.reason,
            row.open_incidents,
            "; ".join(row.discrepancies),
            row.responsible,
            row.date_value,
            row.observations,
        ]
        for column, value in enumerate(values, start=1):
            cell = detail.cell(row=index, column=column, value=value)
            cell.border = BORDER_ALL
            cell.font = Font(name="Aptos", size=9, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 5, 12, 13, 17, 19, 22))
            if column == 21:
                cell.number_format = "DD/MM/YYYY"
        detail.row_dimensions[index].height = 32
    last = 4 + len(results)
    table = Table(displayName="TablaResultado", ref=f"A4:V{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    detail.add_table(table)
    detail.freeze_panes = "A5"
    widths = {
        "A": 18, "B": 10, "C": 32, "D": 26, "E": 22, "F": 10, "G": 20,
        "H": 14, "I": 13, "J": 20, "K": 20, "L": 30, "M": 30, "N": 16,
        "O": 16, "P": 30, "Q": 65, "R": 14, "S": 34, "T": 20, "U": 16, "V": 28,
    }
    for column, width in widths.items():
        detail.column_dimensions[column].width = width
    detail.print_title_rows = "1:4"
    detail.page_setup.orientation = "landscape"
    detail.page_setup.fitToWidth = 1
    detail.sheet_properties.pageSetUpPr.fitToPage = True
    add_detail_formatting(detail, last)

    incidents = wb.create_sheet("Incidencias")
    incidents.sheet_properties.tabColor = "A93226"
    incidents.sheet_view.showGridLines = False
    title_block(incidents, "L", "Incidencias documentadas y discrepancias detectadas en la conciliación")
    report_issues: list[tuple[Issue, str]] = [(issue, "Archivo") for issue in issues]
    documented = {
        (issue.code, issue.incident_type, issue.description)
        for issue in issues
    }
    for row in results:
        if not row.discrepancies:
            continue
        description = "; ".join(row.discrepancies)
        key = (row.code, "Datos maestros", description)
        if key in documented:
            continue
        documented.add(key)
        report_issues.append((
            Issue(
                code=row.code,
                incident_type="Datos maestros",
                description=description,
                evidence="Contraste con el inventario institucional local",
                responsible="",
                date_value=None,
                status="Abierta",
                resolution="",
                observations="Detectada por el procesador; requiere revisión administrativa.",
                blocking=True,
            ),
            "Procesador",
        ))
    incident_headers = INCIDENT_HEADERS + ["Bloqueante", "Error de validación", "Origen"]
    header_row(incidents, 4, incident_headers)
    for index, (issue, origin) in enumerate(report_issues, start=5):
        values = [
            issue.code, issue.incident_type, issue.description, issue.evidence,
            issue.responsible, issue.date_value, issue.status, issue.resolution,
            issue.observations, "Sí" if issue.blocking else "No", "", origin,
        ]
        for column, value in enumerate(values, start=1):
            cell = incidents.cell(row=index, column=column, value=value)
            cell.border = BORDER_ALL
            cell.font = Font(name="Aptos", size=9, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 8, 9))
            if column == 6:
                cell.number_format = "DD/MM/YYYY"
        incidents.row_dimensions[index].height = 30
    if report_issues:
        table = Table(displayName="TablaIncidenciasResultado", ref=f"A4:L{4 + len(report_issues)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        incidents.add_table(table)
    incidents.freeze_panes = "A5"
    for column, width in {"A": 18, "B": 28, "C": 42, "D": 30, "E": 22, "F": 16, "G": 22, "H": 35, "I": 28, "J": 12, "K": 18, "L": 14}.items():
        incidents.column_dimensions[column].width = width

    meta = wb.create_sheet("Metadatos")
    meta.sheet_properties.tabColor = MUTED
    meta.sheet_view.showGridLines = False
    title_block(meta, "B", "Trazabilidad del procesamiento")
    meta_rows = [
        ("Versión del procesador", PROCESSOR_VERSION),
        ("Fecha de procesamiento", processed_at),
        ("Hash SHA-256 del Excel de entrada", source_hash),
        ("Ejemplares de referencia en la base local", reference_count),
        ("Ejecuciones repetidas", "El resultado depende del Excel y de la base; no se escribe en ninguno de los dos."),
        ("Errores globales", "\n".join(global_errors) if global_errors else "Ninguno."),
    ]
    for index, (key, value) in enumerate(meta_rows, start=5):
        meta[f"A{index}"] = key
        meta[f"B{index}"] = value
        meta[f"A{index}"].font = Font(name="Aptos", size=10, bold=True, color=MUTED)
        meta[f"B{index}"].font = Font(name="Aptos", size=10, color=INK)
        meta[f"B{index}"].alignment = Alignment(wrap_text=True, vertical="top")
        meta.row_dimensions[index].height = 26
    meta.column_dimensions["A"].width = 44
    meta.column_dimensions["B"].width = 100

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.title = "Resultado de conciliación física · Trueque IUB"
    wb.properties.creator = "Trueque IUB"
    wb.properties.subject = "Informe de lectura; no modifica inventario"
    wb.calculation.calcMode = "auto"
    wb.save(output)


def add_detail_formatting(ws, last: int) -> None:
    def rule(formula: str, fill: str, color: str) -> FormulaRule:
        return FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=fill), font=Font(color=color, bold=True))

    ws.conditional_formatting.add(f"P5:P{last}", rule("$P5=\"Candidato a habilitación\"", GREEN_SOFT, GREEN))
    ws.conditional_formatting.add(f"P5:P{last}", rule("OR($P5=\"No encontrado\",$P5=\"No habilitable con la información disponible\",$P5=\"Error de validación\",$P5=\"Identificación no coincidente\")", RED_SOFT, RED))
    ws.conditional_formatting.add(f"P5:P{last}", rule("OR($P5=\"Pendiente de conteo\",$P5=\"Pendiente de decisión administrativa\",$P5=\"Estado físico pendiente\",$P5=\"Con discrepancia\")", AMBER_SOFT, AMBER))
    ws.conditional_formatting.add(f"O5:O{last}", rule("$O5=\"No\"", RED_SOFT, RED))
    ws.conditional_formatting.add(f"R5:R{last}", rule("$R5>0", AMBER_SOFT, AMBER))
    ws.conditional_formatting.add(f"S5:S{last}", rule("LEN(TRIM($S5))>0", RED_SOFT, RED))


def process(input_path: Path, url: str, output: Path, processed_at: str) -> dict[str, object]:
    before_hash = hashlib.sha256(input_path.read_bytes()).hexdigest()
    workbook = load_workbook(input_path, data_only=True)
    reference = load_reference(url)
    if len(reference) != 67:
        raise SystemExit(f"Expected 67 reference copies, found {len(reference)}.")
    results, global_errors, issues = read_detail(workbook, reference)
    after_hash = hashlib.sha256(input_path.read_bytes()).hexdigest()
    if before_hash != after_hash:
        raise SystemExit("The input workbook changed during processing.")
    summary = classify_summary(results, global_errors, len(reference))
    summary["incidents"] = len(issues)
    fingerprint = hashlib.sha256(json.dumps({
        "source_hash": before_hash,
        "processor": PROCESSOR_VERSION,
        "rows": [[row.code, row.classification, row.reason, row.errors] for row in results],
        "incidents": [issue.__dict__ for issue in issues],
    }, sort_keys=True, default=str).encode("utf-8")).hexdigest()
    build_report(results, issues, global_errors, len(reference), before_hash, processed_at, output)
    return {
        **summary,
        "source_hash": before_hash,
        "processor_version": PROCESSOR_VERSION,
        "processed_at": processed_at,
        "fingerprint": fingerprint,
        "output": str(output.resolve()),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="docs/inventario/Conciliacion_Fisica_Trueque_IUB.xlsx")
    parser.add_argument("--output", default="docs/inventario/Resultado_Conciliacion_Trueque_IUB.xlsx")
    parser.add_argument("--now", help="Fixed ISO timestamp for reproducible tests.")
    args = parser.parse_args()
    url = local_database_url()
    input_path = Path(args.input)
    if not input_path.exists():
        raise SystemExit(f"Input workbook not found: {input_path}")
    processed_at = args.now or datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    result = process(input_path, url, Path(args.output), processed_at)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
