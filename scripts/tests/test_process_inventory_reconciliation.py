#!/usr/bin/env python3
"""Tests for the read-only physical reconciliation processor.

Run with:

    DATABASE_MIGRATION_URL='postgresql://iub_schema_migrator@127.0.0.1:55432/iub_local' \
        python -m unittest scripts/tests/test_process_inventory_reconciliation.py -v

The tests use the local PostgreSQL reference copy and never write to it.
"""

from __future__ import annotations

import hashlib
import importlib.util
import os
import shutil
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import psycopg
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "scripts" / "process-inventory-reconciliation.py"
TEMPLATE = ROOT / "docs" / "inventario" / "Conciliacion_Fisica_Trueque_IUB.xlsx"
FIXED_TIME = "2026-10-28T00:00:00+00:00"

SPEC = importlib.util.spec_from_file_location("inventory_reconciliation", SCRIPT)
assert SPEC and SPEC.loader
processor = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = processor
SPEC.loader.exec_module(processor)


def database_url() -> str:
    return os.environ.get("DATABASE_MIGRATION_URL", "").strip()


def require_local_database() -> str:
    url = database_url()
    if not url:
        raise unittest.SkipTest("DATABASE_MIGRATION_URL is required for these tests.")
    host = (processor.urlparse(url).hostname or "").lower()
    if host not in processor.LOCAL_HOSTS:
        raise unittest.SkipTest(f"Refusing non-local test host: {host}")
    return url


class ReconciliationTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.url = require_local_database()
        if not TEMPLATE.exists():
            raise unittest.SkipTest(f"Template not found: {TEMPLATE}")
        cls.reference = processor.load_reference(cls.url)
        if len(cls.reference) != 67:
            raise unittest.SkipTest("The local database does not hold the 67 imported copies.")

    def setUp(self) -> None:
        self._tmp = TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    # -- helpers ---------------------------------------------------------

    def workbook_copy(self, name: str = "input.xlsx") -> Path:
        target = self.tmp / name
        shutil.copyfile(TEMPLATE, target)
        return target

    def edit(self, path: Path, changes: list[tuple[str, int, int, object]]) -> None:
        wb = load_workbook(path)
        for sheet, row, column, value in changes:
            wb[sheet].cell(row=row, column=column).value = value
        wb.save(path)

    def run_processor(self, path: Path, output_name: str = "result.xlsx") -> dict:
        output = self.tmp / output_name
        return processor.process(path, self.url, output, FIXED_TIME)

    def detail(self, result_path: Path) -> dict[str, dict]:
        ws = load_workbook(result_path)["Detalle"]
        headers = [ws.cell(row=4, column=column).value for column in range(1, 23)]
        rows = {}
        for row in range(5, ws.max_row + 1):
            values = [ws.cell(row=row, column=column).value for column in range(1, 23)]
            if values[0] is None:
                continue
            rows[values[0]] = dict(zip(headers, values))
        return rows

    def check_valid_row(
        self,
        code: str = "IUB-PAZ-001",
        presence: str = "Sí",
        condition: str = "Bueno",
        identification: str = "Coincide",
        availability: str = "Habilitar",
        responsible: str = "Verificadora de prueba",
        date_value: str = "15/10/2026",
    ) -> list[tuple[str, int, int, object]]:
        row = int(code.rsplit("-", 1)[1]) + 4
        return [
            ("Conciliación física", row, 8, presence),
            ("Conciliación física", row, 9, condition),
            ("Conciliación física", row, 10, identification),
            ("Conciliación física", row, 11, availability),
            ("Conciliación física", row, 14, responsible),
            ("Conciliación física", row, 15, date_value),
        ]

    def incident(
        self,
        code: str,
        status: str,
        resolution: str = "",
        incident_type: str = "Presencia física",
    ) -> list[tuple[str, int, int, object]]:
        return [
            ("Incidencias", 5, 1, code),
            ("Incidencias", 5, 2, incident_type),
            ("Incidencias", 5, 3, "Hallazgo documentado para la prueba."),
            ("Incidencias", 5, 4, "Evidencia de prueba"),
            ("Incidencias", 5, 5, "Verificadora de prueba"),
            ("Incidencias", 5, 6, "15/10/2026"),
            ("Incidencias", 5, 7, status),
            ("Incidencias", 5, 8, resolution),
        ]

    def detail_rows(self, result: dict) -> list[dict]:
        return self.detail(Path(result["output"]))

    # -- reference cases -------------------------------------------------

    def test_original_67_codes_are_processed_as_pending(self) -> None:
        path = self.workbook_copy()
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = self.run_processor(path)
        after = hashlib.sha256(path.read_bytes()).hexdigest()

        self.assertEqual(before, after)
        self.assertEqual(result["total_expected"], 67)
        self.assertEqual(result["processed"], 67)
        self.assertEqual(result["pending"], 67)
        self.assertEqual(result["confirmed"], 0)
        self.assertEqual(result["not_found"], 0)
        self.assertEqual(result["candidates"], 0)
        self.assertEqual(result["validation_errors"], 0)
        self.assertEqual(result["source_hash"], before)

        rows = self.detail_rows(result)
        self.assertEqual(len(rows), 67)
        self.assertEqual(set(rows), {f"IUB-PAZ-{number:03d}" for number in range(1, 68)})
        self.assertTrue(all(row["Clasificación"] == "Pendiente de conteo" for row in rows.values()))

    def test_repeated_titles_are_kept_as_separate_copies(self) -> None:
        path = self.workbook_copy()
        result = self.run_processor(path)
        rows = self.detail_rows(result)
        celestina = [code for code, row in rows.items() if row["Título registrado"] == "La Celestina"]
        mujeres = [code for code, row in rows.items() if row["Título registrado"] == "Las Mujeres del Quijote"]
        self.assertEqual(len(celestina), 2)
        self.assertEqual(len(mujeres), 2)
        self.assertEqual(len(set(celestina + mujeres)), 4)

    def test_duplicate_codes_are_rejected(self) -> None:
        path = self.workbook_copy()
        self.edit(path, [("Conciliación física", 5, 1, "IUB-PAZ-002")])
        result = self.run_processor(path)

        self.assertGreaterEqual(result["validation_errors"], 3)
        self.assertEqual(result["processed"], 67)
        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Error de validación")
        self.assertIn("Código esperado", rows["IUB-PAZ-001"]["Motivo de la decisión"])
        self.assertIn("duplicado", rows["IUB-PAZ-002"]["Motivo de la decisión"])

    def test_unknown_code_is_rejected(self) -> None:
        path = self.workbook_copy()
        self.edit(path, [("Conciliación física", 5, 1, "IUB-PAZ-999")])
        result = self.run_processor(path)

        self.assertGreaterEqual(result["validation_errors"], 2)
        rows = self.detail_rows(result)
        self.assertIn("IUB-PAZ-999", rows)
        self.assertEqual(rows["IUB-PAZ-999"]["Clasificación"], "Error de validación")
        self.assertIn("desconocido", rows["IUB-PAZ-999"]["Motivo de la decisión"])

    def test_empty_rows_are_ignored_without_counting_them(self) -> None:
        path = self.workbook_copy()
        self.edit(path, [("Conciliación física", 72, 1, None), ("Conciliación física", 73, 16, None)])
        result = self.run_processor(path)
        self.assertEqual(result["processed"], 67)
        self.assertEqual(result["validation_errors"], 0)

    def test_row_without_code_is_a_validation_error(self) -> None:
        path = self.workbook_copy()
        self.edit(path, [("Conciliación física", 72, 16, "Fila sin código")])
        result = self.run_processor(path)
        self.assertGreaterEqual(result["validation_errors"], 1)
        rows = self.detail_rows(result)
        self.assertIn("(fila 72)", rows)
        self.assertIn("Código institucional ausente", rows["(fila 72)"]["Motivo de la decisión"])

    def test_not_found_copy(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(presence="No", availability="Bloquear"))
        result = self.run_processor(path)

        self.assertEqual(result["not_found"], 1)
        self.assertEqual(result["pending"], 66)
        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "No encontrado")
        self.assertEqual(rows["IUB-PAZ-001"]["Resultado físico"], "No encontrado")

    def test_identification_mismatch(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(identification="No coincide", availability="Revisar"))
        result = self.run_processor(path)

        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Identificación no coincidente")
        self.assertIn("no coincide", rows["IUB-PAZ-001"]["Motivo de la decisión"].lower())

    def test_condition_pending(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(condition="Pendiente", availability="Revisar"))
        result = self.run_processor(path)

        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Estado físico pendiente")
        self.assertIn("falta evaluar el estado", rows["IUB-PAZ-001"]["Motivo de la decisión"])

    def test_open_incident_blocks_candidacy(self) -> None:
        path = self.workbook_copy()
        changes = self.check_valid_row() + self.incident("IUB-PAZ-001", "Abierta")
        self.edit(path, changes)
        result = self.run_processor(path)

        self.assertEqual(result["candidates"], 0)
        self.assertEqual(result["incidents"], 1)
        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Con discrepancia")
        self.assertEqual(rows["IUB-PAZ-001"]["Incidencias abiertas"], 1)
        self.assertIn("bloqueante", rows["IUB-PAZ-001"]["Motivo de la decisión"])

    def test_resolved_incident_does_not_block(self) -> None:
        path = self.workbook_copy()
        changes = self.check_valid_row() + self.incident(
            "IUB-PAZ-001", "Resuelta", resolution="Verificado por coordinación."
        )
        self.edit(path, changes)
        result = self.run_processor(path)

        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Candidato a habilitación")
        self.assertEqual(rows["IUB-PAZ-001"]["Incidencias abiertas"], 0)
        self.assertIn("aprobación institucional", rows["IUB-PAZ-001"]["Motivo de la decisión"])

    def test_invalid_states_are_reported(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(
            presence="Tal vez", condition="Excelente",
            identification="Parecido", availability="Publicar",
        ))
        result = self.run_processor(path)

        self.assertGreaterEqual(result["validation_errors"], 1)
        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Error de validación")
        reason = rows["IUB-PAZ-001"]["Motivo de la decisión"]
        self.assertIn("Presencia inválida", reason)
        self.assertIn("Estado físico inválido", reason)
        self.assertIn("Identificación inválida", reason)
        self.assertIn("Disponibilidad inválida", reason)

    def test_invalid_date_is_reported(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(date_value="31/02/2026"))
        result = self.run_processor(path)

        self.assertGreaterEqual(result["validation_errors"], 1)
        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Error de validación")
        self.assertIn("fecha", rows["IUB-PAZ-001"]["Motivo de la decisión"].lower())
        meta = load_workbook(Path(result["output"]))["Metadatos"]
        self.assertIn("invalid date", str(meta["B10"].value))

    def test_missing_verification_date_is_reported(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(date_value=None))
        result = self.run_processor(path)

        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Error de validación")
        self.assertIn("sin fecha de verificación", rows["IUB-PAZ-001"]["Motivo de la decisión"])

    def test_missing_responsible_is_reported(self) -> None:
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row(responsible=""))
        result = self.run_processor(path)

        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Error de validación")
        self.assertIn("sin responsable", rows["IUB-PAZ-001"]["Motivo de la decisión"])

    def test_altered_bibliographic_data_is_detected(self) -> None:
        path = self.workbook_copy()
        original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        self.edit(path, self.check_valid_row() + [
            ("Conciliación física", 5, 3, "Título alterado"),
            ("Conciliación física", 5, 4, "Autor alterado"),
        ])
        result = self.run_processor(path)
        self.assertNotEqual(result["source_hash"], original_hash)
        self.assertEqual(result["source_hash"], hashlib.sha256(path.read_bytes()).hexdigest())

        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-001"]["Datos registrados coinciden"], "No")
        self.assertEqual(rows["IUB-PAZ-001"]["Clasificación"], "Con discrepancia")
        self.assertIn("Título del archivo no coincide", rows["IUB-PAZ-001"]["Discrepancias detectadas"])
        incidents = load_workbook(Path(result["output"]))["Incidencias"]
        origins = [incidents.cell(row=row, column=12).value for row in range(5, incidents.max_row + 1)]
        self.assertIn("Procesador", origins)

    def test_incomplete_workbook_reports_missing_code(self) -> None:
        path = self.workbook_copy()
        wb = load_workbook(path)
        ws = wb["Conciliación física"]
        for column in range(1, 17):
            ws.cell(row=71, column=column).value = None
        wb.save(path)

        result = self.run_processor(path)
        self.assertEqual(result["total_expected"], 67)
        self.assertEqual(result["processed"], 66)
        rows = self.detail_rows(result)
        self.assertEqual(rows["IUB-PAZ-067"]["Clasificación"], "Error de validación")
        self.assertIn("Código esperado no aparece", rows["IUB-PAZ-067"]["Motivo de la decisión"])

    def test_wrong_sheet_name_aborts(self) -> None:
        path = self.workbook_copy()
        wb = load_workbook(path)
        wb["Conciliación física"].title = "Otra hoja"
        wb.save(path)
        with self.assertRaises(SystemExit):
            self.run_processor(path)

    def test_wrong_headers_abort(self) -> None:
        path = self.workbook_copy()
        self.edit(path, [("Conciliación física", 4, 1, "Código")])
        with self.assertRaises(SystemExit):
            self.run_processor(path)

    def test_repeated_execution_is_deterministic(self) -> None:
        path = self.workbook_copy()
        first = self.run_processor(path, "first.xlsx")
        second = self.run_processor(path, "second.xlsx")

        self.assertEqual(first["fingerprint"], second["fingerprint"])
        self.assertEqual(first["source_hash"], second["source_hash"])
        for key in (
            "total_expected", "processed", "confirmed", "not_found", "pending",
            "discrepancies", "candidates", "not_habilitable", "validation_errors",
        ):
            self.assertEqual(first[key], second[key], key)
        self.assertEqual(self.detail_rows(first), self.detail_rows(second))

    # -- security and database integrity ---------------------------------

    def test_remote_hosts_are_rejected(self) -> None:
        original = os.environ.get("DATABASE_MIGRATION_URL")
        try:
            for host in (
                "ep-remote.neon.tech",
                "trueque-staging.example.com",
                "10.0.0.5",
                "db.example.org",
            ):
                os.environ["DATABASE_MIGRATION_URL"] = f"postgresql://user:pass@{host}/db"
                with self.subTest(host=host):
                    with self.assertRaises(SystemExit) as raised:
                        processor.local_database_url()
                    self.assertIn("Refusing non-local", str(raised.exception))
        finally:
            if original is None:
                os.environ.pop("DATABASE_MIGRATION_URL", None)
            else:
                os.environ["DATABASE_MIGRATION_URL"] = original

    def test_local_hosts_are_accepted(self) -> None:
        original = os.environ.get("DATABASE_MIGRATION_URL")
        try:
            for host in ("localhost", "127.0.0.1", "[::1]", "postgres"):
                os.environ["DATABASE_MIGRATION_URL"] = f"postgresql://user:pass@{host}/db"
                with self.subTest(host=host):
                    self.assertEqual(
                        processor.local_database_url(),
                        f"postgresql://user:pass@{host}/db",
                    )
        finally:
            if original is None:
                os.environ.pop("DATABASE_MIGRATION_URL", None)
            else:
                os.environ["DATABASE_MIGRATION_URL"] = original

    def test_database_is_not_modified(self) -> None:
        def snapshot() -> tuple:
            with psycopg.connect(self.url) as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT count(*), sum(CASE WHEN availability = 'available' THEN 1 ELSE 0 END) FROM public.copies")
                    copies = cursor.fetchone()
                    cursor.execute("SELECT count(*) FROM public.inventory_movements")
                    movements = cursor.fetchone()[0]
                    cursor.execute("SELECT count(*) FROM public.exchange_requests")
                    requests = cursor.fetchone()[0]
                    cursor.execute("SELECT count(*) FROM public.reservations")
                    reservations = cursor.fetchone()[0]
                    cursor.execute("SELECT count(*) FROM public.audit_events")
                    audit_events = cursor.fetchone()[0]
                    cursor.execute("SELECT inventory_code, availability FROM public.copies ORDER BY inventory_code")
                    states = tuple(cursor.fetchall())
            return copies, movements, requests, reservations, audit_events, states

        before = snapshot()
        path = self.workbook_copy()
        self.edit(path, self.check_valid_row())
        self.run_processor(path)
        after = snapshot()
        self.assertEqual(before, after)

    def test_processor_source_has_no_write_statements(self) -> None:
        source = SCRIPT.read_text(encoding="utf-8")
        for statement in ("INSERT INTO", "UPDATE ", "DELETE FROM", "ALTER TABLE", "CREATE TABLE", "DROP "):
            with self.subTest(statement=statement):
                self.assertNotIn(statement, source)


class PureReconciliationTestCase(unittest.TestCase):
    """Checks that do not require PostgreSQL and never write to the database."""

    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.exists():
            raise unittest.SkipTest(f"Template not found: {TEMPLATE}")
        cls.reference = cls.template_reference()

    @classmethod
    def template_reference(cls) -> dict[str, processor.ReferenceCopy]:
        workbook = load_workbook(TEMPLATE, data_only=True)
        sheet = workbook["Conciliación física"]
        reference: dict[str, processor.ReferenceCopy] = {}
        for index, row in enumerate(range(5, 72), start=1):
            code = processor.clean(sheet.cell(row=row, column=1).value)
            reference[code] = processor.ReferenceCopy(
                code=code,
                source_number=processor.clean(sheet.cell(row=row, column=2).value),
                copy_number=processor.clean(sheet.cell(row=row, column=6).value),
                title=processor.clean(sheet.cell(row=row, column=3).value),
                author=processor.clean(sheet.cell(row=row, column=4).value),
                publisher=processor.clean(sheet.cell(row=row, column=5).value),
                location=processor.clean(sheet.cell(row=row, column=7).value),
                availability="unavailable",
                provenance_type="institutional",
                verified=False,
                bibliographic_id=index,
                source_row=row,
            )
        return reference

    def test_template_is_pending_and_deterministic_without_database(self) -> None:
        workbook = load_workbook(TEMPLATE, data_only=True)
        first, first_errors, first_issues = processor.read_detail(workbook, self.reference)
        second, second_errors, second_issues = processor.read_detail(workbook, self.reference)

        self.assertEqual(first, second)
        self.assertEqual(first_errors, second_errors)
        self.assertEqual(first_issues, second_issues)
        self.assertEqual(len(first), 67)
        self.assertEqual(first_errors, [])
        summary = processor.classify_summary(first, first_errors, len(self.reference))
        self.assertEqual(summary["processed"], 67)
        self.assertEqual(summary["confirmed"], 0)
        self.assertEqual(summary["candidates"], 0)
        self.assertEqual(summary["validation_errors"], 0)
        self.assertTrue(all(row.result_physical == "Pendiente" for row in first))
        self.assertTrue(all(not row.confirmed for row in first))

    def test_invalid_row_is_not_counted_as_confirmed(self) -> None:
        workbook = load_workbook(TEMPLATE, data_only=False)
        sheet = workbook["Conciliación física"]
        sheet.cell(row=5, column=1).value = "IUB-PAZ-999"
        sheet.cell(row=5, column=8).value = "Sí"
        sheet.cell(row=5, column=9).value = "Bueno"
        sheet.cell(row=5, column=10).value = "Coincide"
        sheet.cell(row=5, column=11).value = "Habilitar"
        sheet.cell(row=5, column=14).value = "Verificadora de prueba"
        sheet.cell(row=5, column=15).value = "15/10/2026"

        results, errors, _ = processor.read_detail(workbook, self.reference)
        invalid = next(row for row in results if row.code == "IUB-PAZ-999")
        self.assertFalse(invalid.confirmed)
        self.assertEqual(invalid.result_physical, "No válido")
        self.assertEqual(invalid.classification, "Error de validación")
        self.assertTrue(errors)
        summary = processor.classify_summary(results, errors, len(self.reference))
        self.assertEqual(summary["confirmed"], 0)

    def test_regular_or_bad_condition_is_not_an_automatic_candidate(self) -> None:
        for condition in ("Regular", "Malo"):
            with self.subTest(condition=condition):
                workbook = load_workbook(TEMPLATE, data_only=False)
                sheet = workbook["Conciliación física"]
                sheet.cell(row=5, column=8).value = "Sí"
                sheet.cell(row=5, column=9).value = condition
                sheet.cell(row=5, column=10).value = "Coincide"
                sheet.cell(row=5, column=11).value = "Habilitar"
                sheet.cell(row=5, column=14).value = "Verificadora de prueba"
                sheet.cell(row=5, column=15).value = "15/10/2026"

                results, errors, _ = processor.read_detail(workbook, self.reference)
                row = next(item for item in results if item.code == "IUB-PAZ-001")
                self.assertFalse(row.candidate)
                self.assertEqual(row.classification, "Pendiente de decisión administrativa")
                self.assertIn("D-02", row.reason)
                summary = processor.classify_summary(results, errors, len(self.reference))
                self.assertEqual(summary["candidates"], 0)

    def test_remote_hosts_are_rejected_without_database(self) -> None:
        original = os.environ.get("DATABASE_MIGRATION_URL")
        try:
            for host in (
                "ep-remote.neon.tech",
                "trueque-staging.example.com",
                "10.0.0.5",
                "db.example.org",
            ):
                os.environ["DATABASE_MIGRATION_URL"] = f"postgresql://user:pass@{host}/db"
                with self.subTest(host=host):
                    with self.assertRaises(SystemExit) as raised:
                        processor.local_database_url()
                    self.assertIn("Refusing non-local", str(raised.exception))
        finally:
            if original is None:
                os.environ.pop("DATABASE_MIGRATION_URL", None)
            else:
                os.environ["DATABASE_MIGRATION_URL"] = original

    def test_local_hosts_are_accepted_without_database(self) -> None:
        original = os.environ.get("DATABASE_MIGRATION_URL")
        try:
            for host in ("localhost", "127.0.0.1", "[::1]", "postgres"):
                os.environ["DATABASE_MIGRATION_URL"] = f"postgresql://user:pass@{host}/db"
                with self.subTest(host=host):
                    self.assertEqual(
                        processor.local_database_url(),
                        f"postgresql://user:pass@{host}/db",
                    )
        finally:
            if original is None:
                os.environ.pop("DATABASE_MIGRATION_URL", None)
            else:
                os.environ["DATABASE_MIGRATION_URL"] = original

    def test_processor_source_has_no_write_statements(self) -> None:
        source = SCRIPT.read_text(encoding="utf-8")
        for statement in ("INSERT INTO", "UPDATE ", "DELETE FROM", "ALTER TABLE", "CREATE TABLE", "DROP "):
            with self.subTest(statement=statement):
                self.assertNotIn(statement, source)


if __name__ == "__main__":
    unittest.main()
