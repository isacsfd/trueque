#!/usr/bin/env python3
"""Import the institutional inventory from INVENTARIO TRUEQUE.xlsx.

Default mode is a dry run. Use --apply to write to a LOCAL PostgreSQL database.
The script refuses non-local database hosts and never writes to Neon.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
import psycopg
from psycopg.types.json import Jsonb

DEFAULT_FILE = "INVENTARIO TRUEQUE.xlsx"
SHEET = "Hoja1"
EXPECTED_HEADERS = [
    "BIBLIOTECA ACTUAL",
    "No.",
    "AUTOR",
    "TITULO",
    "EJEMPLAR",
    "EDITORIAL",
]
LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1", "postgres"}


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def isbn_for_source(_: object) -> tuple[None, None]:
    # The source has no ISBN column. Absence is represented as NULL, never as
    # a fabricated number.
    return None, None


def source_is_empty(row: tuple[object, ...]) -> bool:
    # Columns C-F (author, title, copy, publisher) all blank means a filler row.
    return all(clean(value) is None for value in row[2:6])


def normalize_for_match(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def assert_local_database(url: str) -> None:
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    if host not in LOCAL_HOSTS:
        raise SystemExit(
            f"Refusing database host '{host}'. This importer only accepts local PostgreSQL."
        )


def read_rows(path: Path) -> tuple[list[dict[str, object]], bytes, list[str]]:
    data = path.read_bytes()
    digest = hashlib.sha256(data).digest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise SystemExit(f"Sheet {SHEET!r} not found in {path}.")
    sheet = workbook[SHEET]
    rows = list(sheet.iter_rows(min_row=1, values_only=True))
    if len(rows) < 2:
        raise SystemExit("The inventory has no header row.")

    headers = [clean(value) for value in rows[1][:6]]
    if headers != EXPECTED_HEADERS:
        raise SystemExit(f"Unexpected headers: {headers!r}")

    parsed: list[dict[str, object]] = []
    for source_row, row in enumerate(rows[2:], start=3):
        values = list(row[:6]) + [None] * max(0, 6 - len(row))
        record = {
            "source_row": source_row,
            "library": clean(values[0]),
            "number": clean(values[1]),
            "author": clean(values[2]),
            "title": clean(values[3]),
            "copy_number": clean(values[4]),
            "publisher": clean(values[5]),
            "empty": source_is_empty(tuple(values)),
            "raw": {
                "BIBLIOTECA ACTUAL": clean(values[0]),
                "No.": clean(values[1]),
                "AUTOR": clean(values[2]),
                "TITULO": clean(values[3]),
                "EJEMPLAR": clean(values[4]),
                "EDITORIAL": clean(values[5]),
            },
        }
        parsed.append(record)
    return parsed, digest, headers


def validate(record: dict[str, object]) -> str | None:
    if record["empty"]:
        return None
    if not record["title"]:
        return "missing_title"
    if not record["author"]:
        return "missing_author"
    if not record["number"]:
        return "missing_source_number"
    if not record["copy_number"]:
        return "missing_copy_number"
    if not record["library"]:
        return "missing_location"
    if not str(record["number"]).isdigit():
        return "invalid_source_number"
    return None


def apply_import(
    connection: psycopg.Connection,
    rows: list[dict[str, object]],
    digest: bytes,
    filename: str,
    headers: list[str],
) -> dict[str, object]:
    result = {
        "total_rows": len(rows),
        "imported": 0,
        "reused": 0,
        "skipped": 0,
        "rejected": 0,
        "skipped_reasons": {},
        "rejected_reasons": {},
        "bibliographic_records": 0,
        "copies": 0,
    }
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))",
            (f"inventory-import:{digest.hex()}:{SHEET}",),
        )
        cursor.execute(
            """
            INSERT INTO public.inventory_import_batches
              (source_filename, source_sha256, source_sheet, metadata)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (source_sha256, source_sheet) DO UPDATE
              SET source_filename = EXCLUDED.source_filename,
                  metadata = EXCLUDED.metadata,
                  completed_at = NULL
            RETURNING id
            """,
            (
                filename,
                digest,
                SHEET,
                Jsonb({"headers": headers, "mode": "local-institutional-inventory"}),
            ),
        )
        batch_id = cursor.fetchone()[0]

        for record in rows:
            source_row = int(record["source_row"])
            raw = record["raw"]
            status = "skipped"
            reason = None
            copy_id = None
            bibliographic_id = None

            rejected = validate(record)
            if record["empty"]:
                reason = "empty_book_row"
            elif rejected:
                status = "rejected"
                reason = rejected
            else:
                library = str(record["library"])
                publisher = record["publisher"]
                publisher_id = None
                if publisher:
                    normalized_publisher = normalize_for_match(str(publisher))
                    cursor.execute(
                        """
                        INSERT INTO public.publishers (name, normalized_name)
                        VALUES (%s, %s)
                        ON CONFLICT (normalized_name) DO UPDATE SET name = EXCLUDED.name
                        RETURNING id
                        """,
                        (publisher, normalized_publisher),
                    )
                    publisher_id = cursor.fetchone()[0]

                cursor.execute(
                    """
                    SELECT id FROM public.bibliographic_records
                    WHERE isbn_normalized IS NULL
                      AND lower(title) = lower(%s)
                      AND lower(author) = lower(%s)
                      AND publisher_id IS NOT DISTINCT FROM %s
                      AND edition IS NULL
                    ORDER BY id LIMIT 1
                    """,
                    (record["title"], record["author"], publisher_id),
                )
                found = cursor.fetchone()
                if found:
                    bibliographic_id = found[0]
                else:
                    isbn, isbn_normalized = isbn_for_source(record["raw"])
                    cursor.execute(
                        """
                        INSERT INTO public.bibliographic_records
                          (title, author, publisher_id, isbn, isbn_normalized)
                        VALUES (%s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        (record["title"], record["author"], publisher_id, isbn, isbn_normalized),
                    )
                    bibliographic_id = cursor.fetchone()[0]

                inventory_code = f"IUB-PAZ-{int(record['number']):03d}"
                provenance = (
                    f"inventory-import:{batch_id}:row:{source_row}:"
                    f"no:{record['number']}:copy:{record['copy_number']}"
                )
                cursor.execute(
                    "SELECT id, bibliographic_record_id, provenance_reference FROM public.copies WHERE inventory_code = %s",
                    (inventory_code,),
                )
                existing = cursor.fetchone()
                if existing:
                    existing_id, existing_bib, existing_reference = existing
                    if existing_bib == bibliographic_id and existing_reference == provenance:
                        copy_id = existing_id
                        status = "reused"
                    else:
                        status = "rejected"
                        reason = "inventory_code_conflict"
                else:
                    cursor.execute(
                        """
                        INSERT INTO public.copies
                          (bibliographic_record_id, inventory_code, physical_condition,
                           availability, location, provenance_type, provenance_reference)
                        VALUES (%s, %s, 'unassessed', 'unavailable', %s, 'institutional', %s)
                        RETURNING id
                        """,
                        (bibliographic_id, inventory_code, library, provenance),
                    )
                    copy_id = cursor.fetchone()[0]
                    status = "imported"

            if reason:
                bucket = result["skipped_reasons"] if status == "skipped" else result["rejected_reasons"]
                bucket[reason] = bucket.get(reason, 0) + 1
            if status in ("imported", "reused", "skipped", "rejected"):
                result[status] += 1

            cursor.execute(
                """
                INSERT INTO public.inventory_import_rows
                  (batch_id, source_row, source_number, source_copy_number, status,
                   bibliographic_record_id, copy_id, reason, raw_row)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (batch_id, source_row) DO UPDATE
                  SET source_number = EXCLUDED.source_number,
                      source_copy_number = EXCLUDED.source_copy_number,
                      status = EXCLUDED.status,
                      bibliographic_record_id = EXCLUDED.bibliographic_record_id,
                      copy_id = EXCLUDED.copy_id,
                      reason = EXCLUDED.reason,
                      raw_row = EXCLUDED.raw_row,
                      updated_at = clock_timestamp()
                """,
                (
                    batch_id,
                    source_row,
                    record["number"],
                    record["copy_number"],
                    status,
                    bibliographic_id,
                    copy_id,
                    reason,
                    Jsonb(raw),
                ),
            )

        cursor.execute(
            """
            UPDATE public.inventory_import_batches
            SET completed_at = clock_timestamp(),
                total_rows = %s,
                imported_rows = %s,
                reused_rows = %s,
                skipped_rows = %s,
                rejected_rows = %s
            WHERE id = %s
            """,
            (
                result["total_rows"],
                result["imported"],
                result["reused"],
                result["skipped"],
                result["rejected"],
                batch_id,
            ),
        )
        cursor.execute(
            """
            SELECT count(DISTINCT bibliographic_record_id)
            FROM public.inventory_import_rows
            WHERE batch_id = %s AND bibliographic_record_id IS NOT NULL
            """,
            (batch_id,),
        )
        result["bibliographic_records"] = cursor.fetchone()[0]
        cursor.execute(
            """
            SELECT count(*) FROM public.inventory_import_rows
            WHERE batch_id = %s AND copy_id IS NOT NULL
            """,
            (batch_id,),
        )
        result["copies"] = cursor.fetchone()[0]
        result["batch_id"] = batch_id
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", default=os.environ.get("INVENTORY_XLSX", DEFAULT_FILE))
    parser.add_argument("--apply", action="store_true", help="Write to the local database.")
    args = parser.parse_args()

    path = Path(args.file)
    if not path.is_absolute():
        path = Path.cwd() / path
    if not path.exists():
        raise SystemExit(f"Inventory file not found: {path}")

    rows, digest, headers = read_rows(path)
    summary: dict[str, object] = {
        "file": str(path),
        "sha256": digest.hex(),
        "sheet": SHEET,
        "total_rows": len(rows),
        "mode": "apply" if args.apply else "dry-run",
    }
    if not args.apply:
        counts = {"imported": 0, "reused": 0, "skipped": 0, "rejected": 0}
        for record in rows:
            if record["empty"]:
                counts["skipped"] += 1
            elif validate(record):
                counts["rejected"] += 1
            else:
                counts["imported"] += 1
        summary.update(counts)
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 0

    url = os.environ.get("DATABASE_MIGRATION_URL")
    if not url:
        raise SystemExit("DATABASE_MIGRATION_URL is required for --apply.")
    assert_local_database(url)

    with psycopg.connect(url) as connection:
        summary.update(apply_import(connection, rows, digest, path.name, headers))
        connection.commit()
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
