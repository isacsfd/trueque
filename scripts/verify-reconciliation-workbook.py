#!/usr/bin/env python3
"""Verify the generated reconciliation workbook without modifying it."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_SHEETS = ["Conciliación física", "Resumen", "Incidencias", "Instrucciones"]
EXPECTED_CODES = [f"IUB-PAZ-{number:03d}" for number in range(1, 68)]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--file",
        default="docs/inventario/Conciliacion_Fisica_Trueque_IUB.xlsx",
    )
    args = parser.parse_args()
    path = Path(args.file)
    if not path.exists():
        raise SystemExit(f"Workbook not found: {path}")

    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == EXPECTED_SHEETS, wb.sheetnames
    ws = wb["Conciliación física"]
    codes = [ws.cell(row=row, column=1).value for row in range(5, 72)]
    assert len(codes) == 67
    assert len(set(codes)) == 67
    assert codes == EXPECTED_CODES
    for row in range(5, 72):
        for column in range(1, 8):
            assert ws.cell(row=row, column=column).value not in (None, ""), (row, column)
        for column in range(8, 12):
            assert ws.cell(row=row, column=column).value is None, (row, column)
    assert ws.freeze_panes == "A5"
    assert "TablaConciliacion" in ws.tables
    assert ws.tables["TablaConciliacion"].ref == "A4:P71"
    assert len(ws.data_validations.dataValidation) == 5
    assert len(list(ws.conditional_formatting)) >= 5

    summary = wb["Resumen"]
    formulas = {row: summary.cell(row=row, column=2).value for row in range(4, 18)}
    required = {
        "B4": "COUNTA",
        "B5": "COUNTIF",
        "B10": "SUMPRODUCT",
        "B14": "Habilitar",
        "B15": "Bloquear",
        "B16": "COUNTBLANK",
        "B17": "Incidencias",
    }
    for cell, needle in required.items():
        assert needle in str(formulas[int(cell[1:])]), (cell, formulas[int(cell[1:])])
    assert "A5:A71" in formulas[4]
    assert "H5:H71" in formulas[5]
    assert "K5:K71" in formulas[16]
    assert formulas[10].count("(") == formulas[10].count(")")

    incidents = wb["Incidencias"]
    assert "TablaIncidencias" in incidents.tables
    assert incidents.tables["TablaIncidencias"].ref == "A4:I24"
    assert len(incidents.data_validations.dataValidation) == 3

    assert wb.calculation.calcMode == "auto"
    assert wb.calculation.fullCalcOnLoad is True

    print(json.dumps({
        "file": str(path.resolve()),
        "sheets": wb.sheetnames,
        "data_rows": len(codes),
        "unique_codes": len(set(codes)),
        "first_code": codes[0],
        "last_code": codes[-1],
        "summary_formulas": len(formulas),
        "conciliacion_validations": len(ws.data_validations.dataValidation),
        "incidencias_validations": len(incidents.data_validations.dataValidation),
    }, ensure_ascii=False, indent=2))
    print("Workbook verification passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
