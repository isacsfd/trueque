#!/usr/bin/env python3
"""Verify the reconciliation template against the original inventory workbook.

This check is read-only, does not connect to PostgreSQL and does not modify
either workbook. It proves that the 67 institutional codes are unique and that
each row matches the bibliographic and physical data in the source Excel.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_ROWS = 67
EXPECTED_BIBLIOGRAPHIC_RECORDS = 65
SOURCE_HEADER_ROW = 2
TEMPLATE_HEADER_ROW = 4
TEMPLATE_FIRST_ROW = 5
TEMPLATE_LAST_ROW = 71


def normalize(value: object) -> str:
    return " ".join("" if value is None else str(value).split()).casefold()


def source_number(value: object) -> str:
    text = normalize(value)
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError:
        return text
    return str(int(number)) if number.is_integer() else text


def read_source(path: Path) -> tuple[list[dict[str, str]], str]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "Hoja1" not in workbook.sheetnames:
        raise SystemExit(f"La hoja Hoja1 no existe en {path}.")
    sheet = workbook["Hoja1"]
    headers = [normalize(sheet.cell(row=SOURCE_HEADER_ROW, column=column).value) for column in range(1, 7)]
    expected = ["biblioteca actual", "no.", "autor", "titulo", "ejemplar", "editorial"]
    if headers != expected:
        raise SystemExit(f"Encabezados inesperados en el Excel de origen: {headers!r}")

    rows: list[dict[str, str]] = []
    for row in range(SOURCE_HEADER_ROW + 1, sheet.max_row + 1):
        location, number, author, title, copy_number, publisher = [
            sheet.cell(row=row, column=column).value for column in range(1, 7)
        ]
        if not any(value not in (None, "") for value in (title, author, publisher)):
            continue
        rows.append({
            "source_row": str(row),
            "source_number": source_number(number),
            "location": normalize(location),
            "author": normalize(author),
            "title": normalize(title),
            "copy_number": source_number(copy_number),
            "publisher": normalize(publisher),
            "title_display": " ".join(str(title).split()),
        })
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return rows, digest


def read_template(path: Path) -> tuple[list[dict[str, str]], str | None]:
    workbook = load_workbook(path, data_only=False)
    if "Conciliación física" not in workbook.sheetnames:
        raise SystemExit("La plantilla no contiene la hoja Conciliación física.")
    sheet = workbook["Conciliación física"]
    headers = [normalize(sheet.cell(row=TEMPLATE_HEADER_ROW, column=column).value) for column in range(1, 17)]
    expected = [
        "código institucional", "número fuente", "título", "autor", "editorial",
        "número / identificación del ejemplar", "ubicación registrada",
    ]
    if headers[:7] != expected:
        raise SystemExit(f"Encabezados inesperados en la plantilla: {headers[:7]!r}")

    rows: list[dict[str, str]] = []
    for row in range(TEMPLATE_FIRST_ROW, TEMPLATE_LAST_ROW + 1):
        values = [sheet.cell(row=row, column=column).value for column in range(1, 17)]
        if all(value in (None, "") for value in values):
            continue
        rows.append({
            "excel_row": str(row),
            "code": " ".join(str(values[0] or "").split()),
            "source_number": source_number(values[1]),
            "title": normalize(values[2]),
            "author": normalize(values[3]),
            "publisher": normalize(values[4]),
            "copy_number": source_number(values[5]),
            "location": normalize(values[6]),
            "title_display": " ".join(str(values[2] or "").split()),
        })

    summary_hash = None
    if "Resumen" in workbook.sheetnames:
        summary_hash = workbook["Resumen"]["B23"].value
    return rows, summary_hash


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", default="INVENTARIO TRUEQUE.xlsx")
    parser.add_argument(
        "--template",
        default="docs/inventario/Conciliacion_Fisica_Trueque_IUB.xlsx",
    )
    args = parser.parse_args()

    source_path = Path(args.source)
    template_path = Path(args.template)
    if not source_path.exists():
        raise SystemExit(f"Excel de origen no encontrado: {source_path}")
    if not template_path.exists():
        raise SystemExit(f"Plantilla no encontrada: {template_path}")

    source_rows, source_hash = read_source(source_path)
    template_rows, summary_hash = read_template(template_path)
    errors: list[str] = []

    if len(source_rows) != EXPECTED_ROWS:
        errors.append(f"El origen contiene {len(source_rows)} filas con datos; se esperaban {EXPECTED_ROWS}.")
    if len(template_rows) != EXPECTED_ROWS:
        errors.append(f"La plantilla contiene {len(template_rows)} filas; se esperaban {EXPECTED_ROWS}.")

    source_numbers = [row["source_number"] for row in source_rows]
    template_numbers = [row["source_number"] for row in template_rows]
    codes = [row["code"] for row in template_rows]
    if len(set(source_numbers)) != len(source_numbers):
        errors.append("El Excel de origen contiene números fuente duplicados.")
    if len(set(template_numbers)) != len(template_numbers):
        errors.append("La plantilla contiene números fuente duplicados.")
    if len(set(codes)) != len(codes):
        errors.append("La plantilla contiene códigos institucionales duplicados.")
    if set(source_numbers) != set(template_numbers):
        missing = sorted(set(source_numbers) - set(template_numbers))
        extra = sorted(set(template_numbers) - set(source_numbers))
        errors.append(f"Números fuente no conciliados. Faltan: {missing!r}; sobran: {extra!r}.")
    if set(codes) != {f"IUB-PAZ-{number:03d}" for number in range(1, EXPECTED_ROWS + 1)}:
        errors.append("Los códigos institucionales no son IUB-PAZ-001 a IUB-PAZ-067.")

    source_by_number = {row["source_number"]: row for row in source_rows}
    for row in template_rows:
        expected = source_by_number.get(row["source_number"])
        if expected is None:
            continue
        number = int(row["source_number"])
        expected_code = f"IUB-PAZ-{number:03d}"
        if row["code"] != expected_code:
            errors.append(f"Fila {row['excel_row']}: código {row['code']!r}; se esperaba {expected_code!r}.")
        for field, label in (
            ("title", "título"),
            ("author", "autor"),
            ("publisher", "editorial"),
            ("copy_number", "ejemplar"),
            ("location", "ubicación"),
        ):
            if row[field] != expected[field]:
                errors.append(
                    f"Fila {row['excel_row']} ({row['code']}): {label} no coincide "
                    f"({row[field]!r} != {expected[field]!r})."
                )

    bibliographic_keys = {
        (row["title"], row["author"], row["publisher"]) for row in source_rows
    }
    if len(bibliographic_keys) != EXPECTED_BIBLIOGRAPHIC_RECORDS:
        errors.append(
            f"El origen contiene {len(bibliographic_keys)} registros bibliográficos distintos; "
            f"se esperaban {EXPECTED_BIBLIOGRAPHIC_RECORDS}."
        )

    duplicate_titles = {
        title: count
        for title, count in Counter(row["title_display"] for row in source_rows).items()
        if count > 1
    }
    if duplicate_titles != {"La Celestina": 2, "Las Mujeres del Quijote": 2}:
        errors.append(f"Duplicados legítimos inesperados: {duplicate_titles!r}.")

    if normalize(summary_hash) != source_hash:
        errors.append(
            "El hash SHA-256 del origen no coincide con Resumen!B23 de la plantilla."
        )

    result = {
        "source": str(source_path.resolve()),
        "source_sha256": source_hash,
        "source_data_rows": len(source_rows),
        "template": str(template_path.resolve()),
        "template_rows": len(template_rows),
        "unique_codes": len(set(codes)),
        "unique_source_numbers": len(set(source_numbers)),
        "bibliographic_records": len(bibliographic_keys),
        "duplicate_titles": duplicate_titles,
        "hash_matches_template": normalize(summary_hash) == source_hash,
        "errors": errors,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        print("Source reconciliation verification failed.")
        return 1
    print("Source reconciliation verification passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
